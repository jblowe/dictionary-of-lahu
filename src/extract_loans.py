#!/usr/bin/env python3
"""extract_loans.py -- pull every LOAN/LOAN?-tagged entry out of the flat
CSV export and write a formatted spreadsheet of loanwords, with a
best-effort source language extracted from each entry's etymology note.

Source: generated/lahu-flat.csv (built by extract_flat_file.py from
generated/tei/lahu.xml), so this script is a pure filter/reformat step
over already-flattened data -- no XML or band-file parsing of its own.

The source-language extraction (TOKEN_PATTERNS / loan_source_from_notes
below) is the canonical copy of this logic; src/build_search_db.py
carries its own copy (see that file's "Loanword source-language
extraction" section) so the website's search_loan_source column stays
in sync -- if you tune the patterns here, update build_search_db.py's
copy to match, then rebuild the site's database.

Usage:
    python3 src/extract_loans.py [path/to/lahu-flat.csv] [-o output.xlsx]

Run with no arguments from the project root to use the defaults
(generated/lahu-flat.csv -> generated/lahu-loanwords.xlsx). Requires
openpyxl (`pip install openpyxl --break-system-packages` if missing).
"""

import argparse
import csv
import re
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Loanword source-language extraction (canonical copy -- see module
# docstring; src/build_search_db.py duplicates this for the website)
# ---------------------------------------------------------------------------

TOKEN_PATTERNS = [
    (r'N\.?\s*Thai', 'Northern Thai'),
    (r'No\.?\s*Thai', 'Northern Thai'),
    (r'N\.?\s*Tai', 'Northern Thai'),
    (r'Brit\.?\s*Eng\.?', 'English (British)'),
    (r'Written Burmese', 'Burmese (Written)'),
    (r'Written Tibetan', 'Tibetan (Written)'),
    (r'Proto-Lolo-Burmese', 'Proto-Lolo-Burmese (reconstructed)'),
    (r'Proto-Loloish', 'Proto-Loloish (reconstructed)'),
    (r'Proto-Tai', 'Proto-Tai (reconstructed)'),
    (r'Proto-Tibeto-Burman', 'Proto-Tibeto-Burman (reconstructed)'),
    (r'PLB', 'Proto-Lolo-Burmese (reconstructed)'),
    (r'PLoloish', 'Proto-Loloish (reconstructed)'),
    (r'PTB', 'Proto-Tibeto-Burman (reconstructed)'),
    (r'PT\b', 'Proto-Tai (reconstructed)'),
    (r'PL\b', 'Proto-Loloish (reconstructed)'),
    (r'TN\b', 'Tai Nuea'),
    (r'Tai Nuea', 'Tai Nuea'),
    (r'Shan', 'Shan'),
    (r'Siamese', 'Siamese/Thai'),
    (r'Si\.', 'Siamese/Thai'),
    (r'Thai', 'Thai'),
    (r'Tai', 'Tai (general)'),
    (r'Bs\.', 'Burmese'),
    (r'Burmese', 'Burmese'),
    (r'WB\b', 'Written Burmese'),
    (r'WT\b', 'Written Tibetan'),
    (r'Mand\.', 'Mandarin Chinese'),
    (r'Mandarin', 'Mandarin Chinese'),
    (r'Chin\.', 'Chinese'),
    (r'Chinese', 'Chinese'),
    (r'Ch\.', 'Chinese'),
    (r'Eng\.', 'English'),
    (r'English', 'English'),
    (r'Skt\.', 'Sanskrit'),
    (r'Sanskrit', 'Sanskrit'),
    (r'Pali', 'Pali'),
    (r'Jg\.', 'Jingpho'),
    (r'Jingpho', 'Jingpho'),
    (r'Jse\.', 'Japanese'),
    (r'Japanese', 'Japanese'),
    (r'Khmer', 'Khmer'),
    (r'Hindi', 'Hindi'),
    (r'Latin', 'Latin'),
    (r'Greek', 'Greek'),
    (r'Yid\.', 'Yiddish'),
    (r'Fr\.', 'French'),
    (r'French', 'French'),
]
_TOKEN_RE = re.compile('|'.join(f'(?P<t{i}>{pat})' for i, (pat, _) in enumerate(TOKEN_PATTERNS)))
_NAMES = [name for _, name in TOKEN_PATTERNS]
_FILLER_RE = re.compile(
    r'^(?:\s*(?:prob\.?|ult\.?|perhaps|also|the|same|possibly|apparently|first\s+syll\.?|'
    r'1st\.?\s*syll\.?|2nd\.?\s*syll\.?|last\s+element|missionary)\s*)*',
    re.IGNORECASE,
)


def loan_source_from_notes(notes_text):
    """Best-effort extraction of the proximate source language named in an
    etymology note, e.g. "< Shan s'oo" -> 'Shan'. Returns '' if the note
    doesn't name one the abbreviation key recognizes (this does NOT mean
    the word isn't a loan -- Matisoff tagged it LOAN regardless)."""
    if not notes_text:
        return ''
    for m in re.finditer(r'<', notes_text):
        rest = notes_text[m.end():m.end() + 60]
        fm = _FILLER_RE.match(rest)
        rest2 = rest[fm.end():] if fm else rest
        tm = _TOKEN_RE.match(rest2.strip())
        if tm:
            idx = [g for g in tm.groupdict() if tm.group(g)][0]
            return _NAMES[int(idx[1:])]
    return ''


_ILLEGAL_XLSX_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean(s):
    """Strip control characters openpyxl refuses to write to a cell."""
    return _ILLEGAL_XLSX_RE.sub('', s) if isinstance(s, str) else s


def load_loan_rows(flat_csv_path):
    rows = []
    with open(flat_csv_path, encoding='utf-8') as f:
        for row in csv.DictReader(f):
            if not row.get('loan_marker'):
                continue
            display_headword = row['subheadword'] or row['headword']
            under = row['headword'] if row['subheadword'] else ''
            rows.append({
                'headword': display_headword,
                'under': under,
                'part_of_speech': row['part_of_speech'],
                'loan_marker': row['loan_marker'],
                'source_language': loan_source_from_notes(row['notes']) or 'unspecified',
                'definition': row['definition'],
                'notes': row['notes'],
            })
    return rows


def build_workbook(rows):
    wb = Workbook()

    # --- Sheet 1: full list ---
    ws = wb.active
    ws.title = 'Loanwords'
    headers = ['Headword', 'Sub-entry under', 'Part of speech', 'Marker',
               'Source language', 'Definition', 'Etymology note (as in dictionary)']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
        cell.alignment = Alignment(vertical='center')

    for r in rows:
        ws.append([clean(r['headword']), clean(r['under']), clean(r['part_of_speech']),
                   clean(r['loan_marker']), clean(r['source_language']),
                   clean(r['definition']), clean(r['notes'])])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=(cell.column_letter in ('F', 'G')))

    widths = [22, 22, 12, 9, 20, 45, 55]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:G{len(rows) + 1}'

    # --- Sheet 2: summary by language ---
    ws2 = wb.create_sheet('Summary by language')
    ws2.append(['Source language', 'Count', 'Share of tagged loanwords'])
    for cell in ws2[1]:
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')

    c = Counter(r['source_language'] for r in rows)
    total = len(rows)
    order = [k for k, _ in c.most_common() if k != 'unspecified'] + (['unspecified'] if 'unspecified' in c else [])
    for lang in order:
        n = c[lang]
        ws2.append([lang, n, n / total if total else 0])

    for row in ws2.iter_rows(min_row=2, min_col=1, max_col=2):
        for cell in row:
            cell.font = Font(name='Arial', size=10)
    for row in ws2.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.font = Font(name='Arial', size=10)
            cell.number_format = '0.0%'

    r = len(order) + 2
    ws2.cell(row=r, column=1, value='Total entries tagged LOAN or LOAN?').font = Font(name='Arial', bold=True)
    ws2.cell(row=r, column=2, value=total).font = Font(name='Arial', bold=True)

    ws2.column_dimensions['A'].width = 34
    ws2.column_dimensions['B'].width = 10
    ws2.column_dimensions['C'].width = 22

    # --- Notes sheet ---
    ws3 = wb.create_sheet('Notes')
    notes = [
        ('What this is', 'Every headword or sub-headword in the dictionary body '
         '(33 letter-sections) that Matisoff tagged LOAN or LOAN? (uncertain '
         'loan), with the source language extracted from the entry’s '
         'etymology note when the dictionary states one.'),
        ('Source data', 'generated/lahu-flat.csv (src/extract_flat_file.py’s '
         'flattening of generated/tei/lahu.xml, the TEI Lex-0 digital edition).'),
        ('"unspecified" language', 'The dictionary marks these words as loans but '
         'the etymology note either doesn’t name a source (no note at all, or '
         'a note that only cross-references a Lahu synonym/related form), states '
         'an unresolved source ("< ?"), or the source couldn’t be reliably '
         'extracted by this script’s pattern-matching. It does NOT mean "not '
         'a loanword" -- all rows in this list are tagged LOAN in the original.'),
        ('How source language was extracted', 'Matisoff’s etymology notes '
         'generally begin ‘< LANGUAGE form’ (e.g. ‘< Shan s’oo’, '
         '‘< Bs. (WB saw’)’), sometimes after a qualifier like ‘prob.’, '
         '‘ult.’, or ‘1st syll.’. This script matches the first such '
         'derivation against an abbreviation key (Shan, Bs.=Burmese, Chin.=Chinese, '
         'Eng.=English, Si./Siamese=Thai, N. Thai/No. Thai=Northern Thai, '
         'PLB/PTB/PT=Proto-language reconstructions, etc). Chains of ultimate vs. '
         'proximate source (e.g. ‘Shan < Bs. < Chinese’) are collapsed to the '
         'first (most proximate) language named.'),
        ('Caveat', 'This is an automated extraction from free-text etymological '
         'notes written for human readers in 1988, not a structured etymology '
         'database -- treat language attributions as a good-faith reading of '
         'Matisoff’s prose, not a guaranteed-accurate parse. Spot-check '
         'anything you plan to rely on.'),
    ]
    ws3.column_dimensions['A'].width = 26
    ws3.column_dimensions['B'].width = 100
    for label, text in notes:
        ws3.append([label, text])
    for row in ws3.iter_rows():
        for i, cell in enumerate(row):
            cell.font = Font(name='Arial', bold=(i == 0), size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    return wb


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('flat_csv', nargs='?', default='generated/lahu-flat.csv',
                     help='path to the flat CSV (default: generated/lahu-flat.csv)')
    ap.add_argument('-o', '--output', default='generated/lahu-loanwords.xlsx',
                     help='output XLSX path (default: generated/lahu-loanwords.xlsx)')
    args = ap.parse_args()

    rows = load_loan_rows(args.flat_csv)
    if not rows:
        sys.exit(f'no LOAN/LOAN?-tagged rows found in {args.flat_csv}')

    wb = build_workbook(rows)
    wb.save(args.output)

    c = Counter(r['source_language'] for r in rows)
    print(f'Wrote {len(rows)} loanword rows to {args.output}', file=sys.stderr)
    for lang, n in c.most_common(15):
        print(f'  {n:5d}  {lang}', file=sys.stderr)


if __name__ == '__main__':
    main()
